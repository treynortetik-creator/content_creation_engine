"""Export API - Download content in various formats."""
import io
import json
import zipfile
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from app.database import get_db
from app.api.auth import get_current_user_id

router = APIRouter()

# Try to import python-docx for DOCX export
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


async def get_job_outputs(job_id: str, user_id: int) -> dict:
    """Get job with all outputs for export."""
    async with get_db() as db:
        # Get job
        cursor = await db.execute(
            """
            SELECT id, original_filename, target_persona, created_at
            FROM jobs WHERE id = ? AND user_id = ?
            """,
            (job_id, user_id)
        )
        job = await cursor.fetchone()

        if not job:
            raise HTTPException(status_code=404, detail="Job not found")

        # Get outputs
        cursor = await db.execute(
            """
            SELECT content_type, variation_number,
                   step3_final, step2_edited, step1_draft,
                   subject_line, preview_text, send_day, email_type, cta_text
            FROM outputs WHERE job_id = ?
            ORDER BY content_type, variation_number
            """,
            (job_id,)
        )

        outputs = []
        for row in await cursor.fetchall():
            content = row["step3_final"] or row["step2_edited"] or row["step1_draft"] or ""
            output = {
                "content_type": row["content_type"],
                "variation_number": row["variation_number"],
                "content": content,
            }
            # Add email sequence fields if present
            if row["content_type"] == "email_sequence":
                output["subject_line"] = row["subject_line"]
                output["preview_text"] = row["preview_text"]
                output["send_day"] = row["send_day"]
                output["email_type"] = row["email_type"]
                output["cta_text"] = row["cta_text"]

            outputs.append(output)

        return {
            "job_id": job["id"],
            "filename": job["original_filename"],
            "persona": job["target_persona"],
            "created_at": job["created_at"],
            "outputs": outputs,
        }


def format_content_markdown(data: dict) -> str:
    """Format job data as markdown."""
    lines = [
        f"# {data['filename']}",
        f"",
        f"*Generated on {data['created_at']}*",
        f"",
        "---",
        "",
    ]

    # Group by content type
    grouped = {}
    for output in data["outputs"]:
        ct = output["content_type"]
        if ct not in grouped:
            grouped[ct] = []
        grouped[ct].append(output)

    type_titles = {
        "linkedin": "LinkedIn Posts",
        "blog": "Blog Posts",
        "email": "Emails",
        "email_sequence": "Email Sequence",
    }

    for content_type, outputs in grouped.items():
        lines.append(f"## {type_titles.get(content_type, content_type.title())}")
        lines.append("")

        for output in outputs:
            if content_type == "email_sequence":
                lines.append(f"### Email {output['variation_number']} - {output.get('email_type', '').replace('_', ' ').title()} (Day {output.get('send_day', 0)})")
                lines.append("")
                if output.get("subject_line"):
                    lines.append(f"**Subject:** {output['subject_line']}")
                    lines.append("")
                if output.get("preview_text"):
                    lines.append(f"*Preview:* {output['preview_text']}")
                    lines.append("")
                lines.append(output["content"])
                lines.append("")
                if output.get("cta_text"):
                    lines.append(f"**CTA:** {output['cta_text']}")
                    lines.append("")
            else:
                lines.append(f"### Variation {output['variation_number']}")
                lines.append("")
                lines.append(output["content"])
                lines.append("")

        lines.append("---")
        lines.append("")

    return "\n".join(lines)


def create_docx(data: dict) -> io.BytesIO:
    """Create DOCX document from job data."""
    if not DOCX_AVAILABLE:
        raise HTTPException(
            status_code=500,
            detail="DOCX export not available. Install python-docx."
        )

    doc = Document()

    # Title
    title = doc.add_heading(data["filename"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    meta = doc.add_paragraph()
    meta.add_run(f"Generated on {data['created_at']}").italic = True
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Group by content type
    grouped = {}
    for output in data["outputs"]:
        ct = output["content_type"]
        if ct not in grouped:
            grouped[ct] = []
        grouped[ct].append(output)

    type_titles = {
        "linkedin": "LinkedIn Posts",
        "blog": "Blog Posts",
        "email": "Emails",
        "email_sequence": "Email Sequence",
    }

    for content_type, outputs in grouped.items():
        doc.add_heading(type_titles.get(content_type, content_type.title()), 1)

        for output in outputs:
            if content_type == "email_sequence":
                email_title = f"Email {output['variation_number']} - {output.get('email_type', '').replace('_', ' ').title()} (Day {output.get('send_day', 0)})"
                doc.add_heading(email_title, 2)

                if output.get("subject_line"):
                    p = doc.add_paragraph()
                    p.add_run("Subject: ").bold = True
                    p.add_run(output["subject_line"])

                if output.get("preview_text"):
                    p = doc.add_paragraph()
                    p.add_run("Preview: ").italic = True
                    p.add_run(output["preview_text"])

                doc.add_paragraph(output["content"])

                if output.get("cta_text"):
                    p = doc.add_paragraph()
                    p.add_run("CTA: ").bold = True
                    p.add_run(output["cta_text"])
            else:
                doc.add_heading(f"Variation {output['variation_number']}", 2)
                doc.add_paragraph(output["content"])

            doc.add_paragraph()

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/export/{job_id}/markdown")
async def export_markdown(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job outputs as Markdown file."""
    data = await get_job_outputs(job_id, user_id)
    markdown = format_content_markdown(data)

    filename = f"{data['filename'].rsplit('.', 1)[0]}_content.md"

    return StreamingResponse(
        io.BytesIO(markdown.encode("utf-8")),
        media_type="text/markdown",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export/{job_id}/json")
async def export_json(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job outputs as JSON file."""
    data = await get_job_outputs(job_id, user_id)

    filename = f"{data['filename'].rsplit('.', 1)[0]}_content.json"

    return StreamingResponse(
        io.BytesIO(json.dumps(data, indent=2).encode("utf-8")),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export/{job_id}/docx")
async def export_docx(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job outputs as DOCX file."""
    if not DOCX_AVAILABLE:
        raise HTTPException(
            status_code=501,
            detail="DOCX export not available. Install: pip install python-docx"
        )

    data = await get_job_outputs(job_id, user_id)
    docx_buffer = create_docx(data)

    filename = f"{data['filename'].rsplit('.', 1)[0]}_content.docx"

    return StreamingResponse(
        docx_buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export/{job_id}/zip")
async def export_zip(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job outputs as ZIP containing all formats."""
    data = await get_job_outputs(job_id, user_id)

    # Create ZIP in memory
    zip_buffer = io.BytesIO()
    base_name = data["filename"].rsplit(".", 1)[0]

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        # Add Markdown
        markdown = format_content_markdown(data)
        zf.writestr(f"{base_name}_content.md", markdown)

        # Add JSON
        zf.writestr(f"{base_name}_content.json", json.dumps(data, indent=2))

        # Add individual content files
        grouped = {}
        for output in data["outputs"]:
            ct = output["content_type"]
            if ct not in grouped:
                grouped[ct] = []
            grouped[ct].append(output)

        for content_type, outputs in grouped.items():
            folder = f"{base_name}/{content_type}"
            for output in outputs:
                if content_type == "email_sequence":
                    fname = f"email_{output['variation_number']}_day_{output.get('send_day', 0)}.txt"
                    content = f"Subject: {output.get('subject_line', '')}\n\n{output['content']}"
                    if output.get("cta_text"):
                        content += f"\n\n[{output['cta_text']}]"
                else:
                    fname = f"variation_{output['variation_number']}.txt"
                    content = output["content"]

                zf.writestr(f"{folder}/{fname}", content)

        # Add DOCX if available
        if DOCX_AVAILABLE:
            try:
                docx_buffer = create_docx(data)
                zf.writestr(f"{base_name}_content.docx", docx_buffer.read())
            except Exception:
                pass  # Skip DOCX if it fails

    zip_buffer.seek(0)
    filename = f"{base_name}_content.zip"

    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export/formats")
async def get_available_formats():
    """Get list of available export formats."""
    formats = [
        {"id": "markdown", "name": "Markdown", "extension": ".md", "available": True},
        {"id": "json", "name": "JSON", "extension": ".json", "available": True},
        {"id": "docx", "name": "Word Document", "extension": ".docx", "available": DOCX_AVAILABLE},
        {"id": "zip", "name": "ZIP Archive", "extension": ".zip", "available": True},
    ]
    return {"formats": formats}
